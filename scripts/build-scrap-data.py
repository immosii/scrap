#!/usr/bin/env python3
"""Build src/data/factory.json from the three source workbooks.

Matching policy:
- Product type only. Color words and model codes are ignored.
- A scrap row uses the coefficient row of that product type.
- If the same part is listed more than once in the coefficient file,
  the first row is kept. Consumption coefficients are never added together.
- A name that lands on two product types (washer and dishwasher)
  stays out of the rate. Two models of one type stay linked; scrap
  quantity is split by production, not by summing coefficients.
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime, date
from openpyxl import load_workbook

TRANS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
YEH = str.maketrans({"ي": "ی", "ك": "ک", "ة": "ه", "أ": "ا", "إ": "ا"})
COLOR = set(
    """
    سفید آبی سبز قرمز صورتی مشکی مات براق استخوانی پاستیلی سیلور نقره نقره‌ای
    کرم رزگلد اینوکس طلایی طلا هاردستون چرمی
    """.split()
)
STATUS = set("نیمه ساخته رنگ شده نیم ساخت جدید مشترک abs".split())
FAMILY_WORDS = set("mini minim min c70 c60 c2s fr16 h wm dw sky metal minimetal".split())
MODEL_WORDS = set("میراکل روسو russo lm70 lm60 lm70c lm70rc مدل".split())
DROP_TOKEN = COLOR | STATUS | FAMILY_WORDS | MODEL_WORDS | {"با", "ای", "و"}


def norm(s: object) -> str:
    s = str(s or "").translate(TRANS).translate(YEH)
    s = s.replace("\u200c", " ").replace("\xa0", " ").lower()
    s = s.replace("×", "x").replace("*", "x")
    s = re.sub(r"(?<=\d)[./](?=\d)", "d", s)
    s = re.sub(r"[()\[\]\{\}\-–—_:،,\\«»\"+]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def explode(token: str) -> list[str]:
    bits = re.findall(r"[a-z0-9d]+|[\u0600-\u06ff]+", token)
    if not bits:
        bits = [token]
    out: list[str] = []
    for bit in bits:
        cur = bit
        if cur in FAMILY_WORDS or cur in MODEL_WORDS or cur in COLOR or cur in STATUS:
            out.append(cur)
            continue
        changed = True
        while changed and cur:
            changed = False
            for pre in ("minimetal", "minim", "mini", "metal", "c70", "c60", "c2s", "fr16", "sky"):
                if cur.startswith(pre) and len(cur) > len(pre):
                    out.append(pre)
                    cur = cur[len(pre) :]
                    changed = True
                    break
        m = re.match(r"^([^\d]{3,})(\d.*)$", cur)
        if m:
            out.append(m.group(1))
            out.append(m.group(2))
        elif cur:
            out.append(cur)
    return out


def is_model_code(token: str) -> bool:
    if token in {"7d5", "10kg", "8t", "9t", "8v", "7v"}:
        return True
    if re.fullmatch(r"\d{1,2}[tv][iw]?\d?", token):
        return True
    if re.fullmatch(r"\d{1,2}kg", token):
        return True
    if re.fullmatch(r"(bl|gr|bk|wh|re|pi)\d*", token):
        return True
    if re.fullmatch(r"0\d", token):
        return True
    return False


def signature(n: str) -> tuple[str, ...]:
    tokens: list[str] = []
    for raw in n.split():
        for token in explode(raw):
            if token in DROP_TOKEN or is_model_code(token) or len(token) <= 1:
                continue
            tokens.append(token)
    return tuple(sorted(set(tokens)))


def dims(sig: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(token for token in sig if re.search(r"\d", token))


def num(x: object) -> float | None:
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).translate(TRANS).strip().replace("٫", ".").replace("/", ".").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def jalali_month(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}/{value.month:02d}"
    s = str(value or "").translate(TRANS).strip()
    m = re.match(r"(\d{4})[/-](\d{1,2})", s)
    if not m:
        return ""
    return f"{int(m.group(1)):04d}/{int(m.group(2)):02d}"


def jalali_date(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}/{value.month:02d}/{value.day:02d}"
    s = str(value or "").translate(TRANS).strip().replace("-", "/")
    m = re.match(r"(\d{4})/(\d{1,2})/(\d{1,2})", s)
    if not m:
        return jalali_month(value)
    return f"{int(m.group(1)):04d}/{int(m.group(2)):02d}/{int(m.group(3)):02d}"


def parse_sheet(ws, mode: str) -> list[dict]:
    rows: list[dict] = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        rr = list(r)
        if mode == "c":
            title, unit, qty, imp = (rr + [None] * 6)[2:6]
        else:
            title, unit, qty, imp = (rr + [None] * 4)[:4]
        if not title:
            continue
        t = str(title).strip()
        if t.startswith("مجموع"):
            continue
        q, im = num(qty), num(imp)
        if q is None or im is None:
            continue
        n = norm(t)
        rows.append(
            {
                "title": t,
                "n": n,
                "ct": list(signature(n)),
                "unit": str(unit or "").strip(),
                "c": q,
                "s": float(im),
                "dups": 1,
            }
        )
    return rows


def parse_bom(path: str) -> list[dict]:
    """Coefficient workbook only. Foam injection belongs to mini, not its own line."""
    wb = load_workbook(path, data_only=True, read_only=True)
    specs = [
        ("wm10", "wm-10", "c"),
        ("wm8a", "wm-8t2-9t1", "p"),
        ("wm8b", "wm-8t3-9t2", "c"),
        ("dw", "DW", "p"),
        ("mini", "MINI", "p"),
        ("mini", "فوم", "p"),
        ("c70", "C70", "p"),
    ]
    bom: list[dict] = []
    seq: dict[str, int] = defaultdict(int)
    seen: dict[tuple, dict] = {}
    for fam, sheet, mode in specs:
        if sheet not in wb.sheetnames:
            continue
        for p in parse_sheet(wb[sheet], mode):
            key = (fam, tuple(p["ct"]) if p["ct"] else ("raw", p["n"]))
            prev = seen.get(key)
            if prev is not None:
                prev["dups"] += 1
                continue
            p["family"] = fam
            p["id"] = f"{fam}-{seq[fam]}"
            seq[fam] += 1
            seen[key] = p
            bom.append(p)
    wb.close()
    return bom


def prefer(sn: str, fams: list[str]) -> list[str]:
    if "rm5" in sn:
        return []
    if "c70" in sn or "fr16" in sn:
        return [f for f in fams if f == "c70"]
    if "mini" in sn or "minim" in sn:
        return [f for f in fams if f == "mini"]
    if any(x in sn for x in ["8t3", "9t2", "8ti3", "9ti2", "8tw3", "9tw2"]):
        return [f for f in fams if f == "wm8b"]
    if any(x in sn for x in ["8t2", "9t1", "8ti2", "9ti1", "8tw2", "9tw1"]):
        return [f for f in fams if f == "wm8a"]
    if "10kg" in sn:
        return [f for f in fams if f == "wm10"]
    if any(x in sn for x in ["8t", "9t", "8v", "7d5"]):
        return [f for f in fams if f in {"wm8a", "wm8b"}]
    if "ظرفشو" in sn or "h dw" in sn or sn.startswith("dw ") or " dw " in f" {sn} ":
        return [f for f in fams if f == "dw"]
    return fams


def product_group(family: str) -> str:
    if family in {"mini", "c70", "rm5"}:
        return "fridge"
    if family == "dw":
        return "dw"
    return "wm"


def match_title(title: str, wh: str, bom: list[dict]) -> tuple[str, list[dict], str]:
    sn = norm(title)
    allow = {"mini", "c70"} if "یخچال" in (wh or "") else {"wm10", "wm8a", "wm8b", "dw"}
    if "rm5" in sn:
        allow = set()
    ct = signature(sn)
    hits: list[tuple[int, dict]] = []
    for p in bom:
        if p["family"] not in allow:
            continue
        pt = tuple(p["ct"])
        ct_use, pt_use = ct, pt
        if "بالانسور" in ct or "بالانسور" in pt:
            # 13 / 15 / 16 on a dishwasher balancer is the model, not another part.
            ct_use = tuple(token for token in ct if token not in {"13", "15", "16"})
            pt_use = tuple(token for token in pt if token not in {"13", "15", "16"})
        if sn == p["n"] or (ct_use and ct_use == pt_use):
            hits.append((100 if sn == p["n"] else 90, p))
            continue
        if not ct_use or not pt_use:
            continue
        if dims(ct_use) and dims(pt_use) and dims(ct_use) != dims(pt_use):
            continue
        a, b = set(ct_use), set(pt_use)
        if a <= b or b <= a:
            short, long_ = (a, b) if len(a) <= len(b) else (b, a)
            extra = long_ - short
            if any(re.search(r"\d", token) for token in extra):
                continue
            strong = len(short) >= 2 or (len(short) == 1 and min(len(token) for token in short) >= 4)
            if strong and extra and all(len(token) <= 2 for token in extra):
                hits.append((80, p))
    if not any(score >= 90 for score, _ in hits):
        subset_hits = []
        for p in bom:
            if p["family"] not in allow:
                continue
            pt = tuple(p["ct"])
            if not ct or not pt or not set(ct) < set(pt):
                continue
            if len(set(pt) - set(ct)) > 2:
                continue
            if dims(ct) and dims(ct) != dims(pt):
                continue
            if not (len(ct) >= 2 or min(len(token) for token in ct) >= 4):
                continue
            subset_hits.append(p)
        by_family_sub: dict[str, list[dict]] = {}
        for p in subset_hits:
            by_family_sub.setdefault(p["family"], []).append(p)
        for parts in by_family_sub.values():
            if len(parts) == 1:
                hits.append((70, parts[0]))
    if not hits:
        return "none", [], ""
    top = max(score for score, _ in hits)
    tops = [p for score, p in hits if score == top]
    by_family: dict[str, dict] = {}
    for p in tops:
        by_family.setdefault(p["family"], p)
    fams = prefer(sn, list(by_family))
    chosen = [by_family[f] for f in fams if f in by_family]
    if not chosen:
        return "none", [], "پیشوند خانواده با BOM مقصد نخواند"
    groups = {product_group(p["family"]) for p in chosen}
    if len(groups) != 1:
        return "ambiguous", [], "نام در چند نوع محصول هست: " + "، ".join(sorted(groups))
    kind = "exact" if top == 100 and all(sn == p["n"] for p in chosen) else "canon"
    if len(chosen) == 1:
        note = "" if kind == "exact" else "فارغ از رنگ و مدل؛ ضریب همین نوع محصول"
        return kind, chosen, note
    return kind, chosen, "همان قطعه در چند مدل این نوع محصول؛ ضریب‌ها جمع نشده و به نسبت تولید تقسیم شده"


def scrap_color(title: str) -> str | None:
    if "صورتی" in title:
        return "pink"
    if "آبی" in title or "Bl-01" in title or "BL-01" in title:
        return "blue"
    if "سبز" in title or "Gr-01" in title:
        return "green"
    if "قرمز" in title:
        return "red"
    if "استخوانی" in title or "سفید" in title:
        return "white"
    if "مات" in title or "BK02" in title or "BK2" in title:
        return "matte"
    if "مشکی" in title or "براق" in title:
        return "gloss"
    if "استیل" in title:
        return "steel"
    if "اینوکس" in title:
        return "inox"
    if "طرح" in title:
        return "pattern"
    return None


def classify_product(name: str, code: str) -> str:
    u = name.upper().replace(" ", "")
    if "ظرفشو" in name or code.startswith("5222"):
        return "dw"
    if "لباس" in name or code.startswith("5221"):
        if "H10" in u or "CWM" in u or "10کیلو" in name.replace(" ", ""):
            return "wm10"
        if any(x in u for x in ["8TI3", "8TW3", "9TI2", "9TW2"]):
            return "wm8b"
        if any(x in u for x in ["8TI2", "8TW2", "9TI1", "9TW1"]):
            return "wm8a"
        return "wm8b"
    if "RM5" in u:
        return "rm5"
    return "mini"


def color_of(name: str) -> str | None:
    u = name.upper()
    if "صورتی" in name or "PI1" in u or "PI3" in u:
        return "pink"
    if "آبی" in name:
        return "blue"
    if "سبز" in name:
        return "green"
    if "قرمز" in name:
        return "red"
    if "استیل" in name or "ST1" in u:
        return "steel"
    if "اینوکس" in name or re.search(r"TI\d", u):
        return "inox"
    if "مات" in name or "BK2" in u or re.search(r"B2\b", u):
        return "matte"
    if "براق" in name or "BK1" in u or "BK3" in u or re.search(r"B1\b", u):
        return "gloss"
    if "سفید" in name or "WH" in u or re.search(r"TW\d|VW\d", u):
        return "white"
    return None


def origin_of(prod: float, con: float) -> str:
    if prod > 0 and con > 0:
        return "both"
    if prod > 0:
        return "prod"
    if con > 0:
        return "contractor"
    return "unknown"


def main() -> None:
    root = "/workspace/attachments"
    bom = parse_bom(f"{root}/شدت اهمیت - ضریب مصرف -2.xlsx")
    collapsed = [p for p in bom if p["dups"] > 1]
    print("kept first row, not summed:")
    for p in collapsed:
        print(f"  {p['family']} x{p['dups']} c={p['c']} s={p['s']} {p['title'][:70]}")

    wb = load_workbook(f"{root}/ضایعات (2).xlsx", data_only=True, read_only=True)
    scrap_rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]
    wb.close()
    scrap = []
    for i, r in enumerate(scrap_rows):
        if not r or r[3] is None:
            continue
        title = str(r[3]).strip()
        wh = str(r[0] or "").strip()
        kind, parts, note = match_title(title, wh, bom)
        pq, cq = num(r[5]) or 0, num(r[6]) or 0
        scrap.append(
            {
                "id": i + 1,
                "month": jalali_month(r[1]),
                "date": jalali_date(r[1]),
                "wh": wh,
                "code": "" if r[2] is None else str(r[2]).split(".")[0].strip(),
                "title": title,
                "reason": "" if r[4] is None else str(r[4]).strip(),
                "origin": origin_of(pq, cq),
                "qty": num(r[7]) or 0,
                "match": kind,
                "bomIds": [p["id"] for p in parts] if kind in ("exact", "canon") else [],
                "color": scrap_color(title),
                "note": note,
            }
        )

    wb = load_workbook(f"{root}/کل تولیدات 6 ماهه اول (2).xlsx", data_only=True, read_only=True)
    prod_rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]
    wb.close()
    models: dict[str, dict] = {}
    cells: dict[tuple, int] = defaultdict(int)
    serials = []
    for r in prod_rows:
        if not r or r[3] is None:
            continue
        code = str(r[3]).split(".")[0].strip()
        name = str(r[2] or "").strip()
        grade = str(r[0] or "A").strip().upper()
        if grade not in ("A", "B"):
            grade = "A"
        month = jalali_month(r[5])
        if code not in models:
            models[code] = {
                "code": code,
                "name": name,
                "family": classify_product(name, code),
                "color": color_of(name),
                "pattern": "طرح" in name,
            }
        cells[(code, month, grade)] += 1
        serials.append(
            {
                "grade": grade,
                "serial": str(r[1] or "").strip(),
                "code": code,
                "month": month,
            }
        )

    bom_out = [
        {
            "id": p["id"],
            "family": p["family"],
            "title": p["title"],
            "unit": p["unit"],
            "c": round(p["c"], 6),
            "s": p["s"],
            "dups": p["dups"],
        }
        for p in bom
    ]
    issues = [
        {
            "level": "warn",
            "title": "بعضی ردیف‌های ضایعات در درصد نمی‌آیند",
            "detail": "فقط وقتی نام کالا — بدون رنگ و مدل — با یک نوع محصول در فایل ضریب یکی باشد وارد درصد می‌شود. اگر در فایل ضریب نباشد، یا هم‌زمان به دو نوع محصول (مثلاً لباسشویی و ظرفشویی) بخورد، خارج از درصد می‌ماند.",
        },
        {
            "level": "info",
            "title": "ردیف تکراری ضریب مصرف جمع نمی‌شود",
            "detail": "اگر یک قطعه چند بار در فایل ضریب آمده باشد، فقط ردیف اول همان نوع محصول استفاده می‌شود. تعداد و درجه اهمیت ردیف‌های بعدی با آن جمع نمی‌شود.",
        },
        {
            "level": "warn",
            "title": "تولید C70 در فایل شش‌ماهه نیست",
            "detail": "برای این خانواده تعداد تولید صفر است، پس درصد ضایعات تعریف نمی‌شود. اگر همان قطعه در مینی هم باشد، مقدار ضایعات با ضریب مینی و تولید مینی حساب می‌شود.",
        },
        {
            "level": "warn",
            "title": "یخچال CM-RM5 فایل ضریب جدا ندارد",
            "detail": "این مدل‌ها در تولید هستند. ضایعاتی که در نامشان RM5 است به شیت مینی یا C70 وصل نمی‌شود.",
        },
        {
            "level": "info",
            "title": "شیت تزریق فوم جزو یخچال مینی است",
            "detail": "قطعات شیت فوم با همان تعداد تولید مینی در مخرج می‌آیند. اگر نامشان با یک قطعه شیت مینی یکی باشد، فقط ردیف اول می‌ماند.",
        },
    ]

    out = {
        "months": ["1405/01", "1405/02", "1405/03", "1405/04", "1405/05", "1405/06"],
        "models": list(models.values()),
        "bom": bom_out,
        "scrap": scrap,
        "prod": [
            {"code": c, "month": m, "grade": g, "qty": n}
            for (c, m, g), n in sorted(cells.items())
        ],
        "serials": serials,
        "issues": issues,
    }
    dest = "/workspace/src/data/factory.json"
    with open(dest, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print("wrote", dest)
    from collections import Counter

    q = sum(s["qty"] for s in scrap)
    qm = sum(s["qty"] for s in scrap if s["match"] in ("exact", "canon"))
    print("models", len(models), "bom", len(bom_out), "scrap", len(scrap), "serials", len(serials))
    print("qty", q, "matched qty", qm, "cover", round(100 * qm / q, 2))
    print("match", Counter(s["match"] for s in scrap))
    print("origin", Counter(s["origin"] for s in scrap))
    print("families", Counter(m["family"] for m in models.values()))
    print("prod", sum(c["qty"] for c in out["prod"]), "B", sum(c["qty"] for c in out["prod"] if c["grade"] == "B"))
    print("--- canon samples ---")
    shown = set()
    for s in scrap:
        if s["match"] != "canon" or s["title"] in shown:
            continue
        shown.add(s["title"])
        ids = s["bomIds"]
        names = [p["title"] for p in bom if p["id"] in ids]
        print(f"  {s['qty']:>6} {s['title'][:60]}")
        print(f"         -> {names}")
        if len(shown) >= 40:
            break


if __name__ == "__main__":
    main()
